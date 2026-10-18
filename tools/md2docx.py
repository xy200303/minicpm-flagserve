"""Render docs/report-draft.md into a Word document and a PDF.

Usage: python tools/md2docx.py <input.md> <output.docx> [output.pdf]
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CJK = "\u5b8b\u4f53"          # SimSun
CJK_HEAD = "\u5fae\u8f6f\u96c5\u9ed1"  # Microsoft YaHei
MONO = "Consolas"


def set_cjk(run, font=CJK):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_runs(par, text, size=None, mono=False, bold=False):
    """Handle **bold** and `code` inline markup (code may nest in bold)."""
    for tok in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_runs(par, tok[2:-2], size, mono, True)
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1].strip("`"))
            r.font.name = MONO
            r.bold = bold
            if size:
                r.font.size = Pt(size - 1)
            set_cjk(r, MONO)
            continue
        r = par.add_run(tok)
        r.bold = bold
        if mono:
            r.font.name = MONO
        if size:
            r.font.size = Pt(size)
        set_cjk(r, MONO if mono else CJK)


def flush_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.space_before = Pt(2)
            txt = row[j] if j < len(row) else ""
            add_runs(par, txt, size=9)
            for r in par.runs:
                r.font.size = Pt(9)
            if i == 0:
                shade(cell, "E8EEF7")
                for r in par.runs:
                    r.bold = True
    doc.add_paragraph()


def main(src, out_docx, out_pdf=None):
    lines = Path(src).read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = CJK
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(56)
        section.top_margin = section.bottom_margin = Pt(56)

    i = 0
    table_rows = []
    in_code = False
    code_buf = []
    para_buf = []
    last_par = {"p": None}

    def close_table():
        nonlocal table_rows
        if table_rows:
            flush_table(doc, table_rows)
            table_rows = []

    def remember(par):
        last_par["p"] = par

    def append_continuation(text):
        par = last_par["p"]
        if par is None:
            para_buf.append(text)
            return
        add_runs(par, " " + text)

    def flush_para():
        nonlocal para_buf
        if not para_buf:
            return
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(5)
        add_runs(par, "".join(para_buf))
        para_buf = []
    last_par = {"p": None}

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        flush_para()
        if stripped.startswith("```"):
            if in_code:
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Pt(12)
                par.paragraph_format.space_after = Pt(6)
                r = par.add_run("\n".join(code_buf))
                r.font.name = MONO
                r.font.size = Pt(8.5)
                set_cjk(r, MONO)
                code_buf = []
                in_code = False
            else:
                close_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        flush_para()
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                i += 1
                continue
            table_rows.append(cells)
            i += 1
            continue
        close_table()

        if not stripped:
            flush_para()
            i += 1
            continue
        if stripped.startswith("---"):
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            flush_para()
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = par.add_run(text)
                r.bold = True
                r.font.size = Pt(20)
                set_cjk(r, CJK_HEAD)
                par.paragraph_format.space_after = Pt(10)
            else:
                h = doc.add_heading(level=min(level, 4), text="")
                r = h.add_run(text)
                r.font.size = Pt(15 if level == 2 else 12.5)
                r.bold = True
                r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
                set_cjk(r, CJK_HEAD)
            i += 1
            continue

        if stripped.startswith("- "):
            flush_para()
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, stripped[2:])
            remember(par)
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            flush_para()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            par.paragraph_format.first_line_indent = Pt(-12)
            par.paragraph_format.space_after = Pt(3)
            add_runs(par, f"{m.group(1)}.  {m.group(2)}")
            remember(par)
            i += 1
            continue
        if stripped.startswith("> "):
            flush_para()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            add_runs(par, stripped[2:], size=9.5)
            for r in par.runs:
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        if line[:1] in (" ", "	") and last_par["p"] is not None and not para_buf:
            append_continuation(stripped)
        else:
            para_buf.append(stripped)
        i += 1

    flush_para()
    close_table()
    if in_code and code_buf:
        par = doc.add_paragraph()
        r = par.add_run("\n".join(code_buf))
        r.font.name = MONO
        r.font.size = Pt(8.5)

    Path(out_docx).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)
    print(f"wrote {out_docx}")

    if out_pdf:
        ps = (
            "$w = New-Object -ComObject Word.Application; $w.Visible = $false;"
            f"$d = $w.Documents.Open('{Path(out_docx).resolve()}');"
            "$d.ExportAsFixedFormat("
            f"'{Path(out_pdf).resolve()}', 17); $d.Close(0); $w.Quit()"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
        print(f"wrote {out_pdf}")


if __name__ == "__main__":
    main(*sys.argv[1:])
